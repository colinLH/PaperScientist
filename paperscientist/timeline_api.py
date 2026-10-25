"""Resumable SSE and read-only document previews for the conversation timeline."""
from __future__ import annotations

import asyncio
import csv
import io
import json
import threading
import time
from pathlib import Path
from urllib.parse import quote, urlparse

from fastapi import HTTPException, Request
from fastapi.responses import FileResponse, Response, StreamingResponse
from pydantic import BaseModel

from .timeline import markdown

PDF_LOCK = threading.Lock()


class UndoRequest(BaseModel):
    turn_id: str
    paths: list[str] | None = None


def pdf_page(data, page):
    import pypdfium2 as pdfium
    with PDF_LOCK:
        with pdfium.PdfDocument(data) as pdf:
            if not 0 <= page < len(pdf):
                raise ValueError('PDF 页码超出范围')
            p = pdf[page]
            bitmap = p.render(scale=min(1.6, 1800 / max(p.get_size())))
            output = io.BytesIO()
            bitmap.to_pil().save(output, format='PNG')
            bitmap.close()
            p.close()
            return output.getvalue()


def preview(data, path, store, sid):
    suffix = Path(path).suffix.lower()
    if suffix == '.pdf':
        import pypdfium2 as pdfium
        with PDF_LOCK:
            with pdfium.PdfDocument(data) as pdf:
                return {'kind': 'pdf', 'pages': len(pdf)}
    if suffix in {'.png', '.jpg', '.jpeg', '.gif', '.webp'}:
        return {'kind': 'image'}
    if suffix == '.docx':
        from docx import Document
        document = Document(io.BytesIO(data))
        return {'kind': 'text', 'text': '\n'.join(p.text for p in document.paragraphs)[:250000]}
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = book.active
        rows = [[str(v) if v is not None else '' for v in row] for row in sheet.iter_rows(max_row=101, max_col=min(sheet.max_column, 30), values_only=True)]
        book.close()
        return {'kind': 'table', 'rows': rows, 'note': '预览首张工作表，最多 100 行、30 列'}
    try:
        text = data[:500000].decode('utf-8-sig')
        if '\0' in text:
            raise UnicodeError()
    except UnicodeError:
        return {'kind': 'download', 'note': '此格式可下载后打开'}
    note = '预览已截取前 500 KB，下载可查看完整文件' if len(data) > 500000 else ''
    if suffix in {'.csv', '.tsv'}:
        reader = csv.reader(io.StringIO(text), delimiter='\t' if suffix == '.tsv' else ',')
        rows = []
        for row in reader:
            rows.append(row[:30])
            if len(rows) == 101:
                break
        return {'kind': 'table', 'rows': rows, 'note': '预览最多 100 行、30 列'}
    if suffix in {'.md', '.markdown'}:
        return {'kind': 'markdown', 'html': markdown(text, store, sid), 'note': note}
    return {'kind': 'text', 'text': text, 'note': note}


def mount_timeline_api(app, service):
    store, timeline, versions = service.store, service.timeline, service.versions

    def require(sid):
        session = store.get_session(sid)
        if not session or session['user_id'] != 'local' or session.get('parent_id'):
            raise HTTPException(404, '会话不存在')
        return session

    def data_for(sid, turn, path):
        require(sid)
        try:
            store.file(sid, path)  # Validate even for an immutable stored version.
            return versions.version_data(sid, turn, path)
        except (ValueError, OSError) as exc:
            raise HTTPException(404, str(exc)) from exc

    @app.get('/api/ui/{asset}')
    async def asset(asset: str):
        if asset not in {'timeline.js', 'timeline.css', 'projects.js', 'projects.css'}:
            raise HTTPException(404)
        return FileResponse(Path(__file__).parent / 'static' / asset, headers={'Cache-Control': 'no-cache'})

    @app.get('/api/sessions/{sid}/timeline')
    async def get_timeline(sid: str):
        require(sid)
        return await asyncio.to_thread(timeline.snapshot, sid)

    @app.get('/api/sessions/{sid}/events')
    async def events(sid: str, request: Request, after: int = 0):
        require(sid)
        try:
            cursor = max(0, after, int(request.headers.get('last-event-id', '0')))
        except ValueError:
            raise HTTPException(400, '无效的事件游标')
        async def stream():
            nonlocal cursor
            last_heartbeat = time.monotonic()
            wake = asyncio.Event()
            loop = asyncio.get_running_loop()
            unsubscribe = timeline.subscribe(sid, lambda: loop.call_soon_threadsafe(wake.set))
            try:
                yield 'retry: 1000\n\n'
                while not await request.is_disconnected():
                    # Clear before reading: a commit during the read/wait boundary
                    # keeps the wakeup set, so no event can be stranded until timeout.
                    wake.clear()
                    rows = await asyncio.to_thread(timeline.after, sid, cursor)
                    for row in rows:
                        cursor = row['id']
                        yield 'id: '+str(cursor)+'\nevent: timeline\ndata: '+json.dumps(row, ensure_ascii=False)+'\n\n'
                    if time.monotonic() - last_heartbeat > 15:
                        if not store.get_session(sid):
                            yield 'event: removed\ndata: {}\n\n'
                            return
                        yield ': keep-alive\n\n'
                        last_heartbeat = time.monotonic()
                    if len(rows) >= 300:
                        continue
                    try:
                        # Fallback also observes writes from a second server process.
                        await asyncio.wait_for(wake.wait(), timeout=2)
                    except asyncio.TimeoutError:
                        pass
            finally:
                unsubscribe()
        return StreamingResponse(stream(), media_type='text/event-stream', headers={'Cache-Control': 'no-cache, no-transform', 'X-Accel-Buffering': 'no'})

    @app.get('/api/sessions/{sid}/tool-detail')
    async def detail(sid: str, event_id: int = 0, message_id: int = 0):
        require(sid)
        with store.connect() as db:
            if event_id:
                row = db.execute('SELECT data FROM events WHERE session_id=? AND id=?', (sid, event_id)).fetchone()
            else:
                row = db.execute('SELECT payload FROM messages WHERE session_id=? AND id=?', (sid, message_id)).fetchone()
        if not row:
            raise HTTPException(404, '工具记录不存在')
        obj = json.loads(row[0])
        if not event_id:
            obj = obj.get('data', {}).get('content', obj)
        text = obj if isinstance(obj, str) else json.dumps(obj, indent=2, ensure_ascii=False)
        return {'text': text[:200000], 'truncated': len(text) > 200000}

    @app.get('/api/sessions/{sid}/file-preview')
    async def file_preview(sid: str, path: str, turn: str = ''):
        data, version = await asyncio.to_thread(data_for, sid, turn, path)
        try:
            result = await asyncio.to_thread(preview, data, path, store, sid)
        except Exception as exc:
            result = {'kind': 'download', 'note': '无法解析预览：'+service.safe_error(exc)}
        return {**result, 'path': path, 'version': version, 'bytes': len(data)}

    @app.get('/api/sessions/{sid}/file-content')
    async def file_content(sid: str, path: str, turn: str = '', page: int = -1):
        data, _ = await asyncio.to_thread(data_for, sid, turn, path)
        if page >= 0:
            if Path(path).suffix.lower() != '.pdf':
                raise HTTPException(400, '仅 PDF 支持分页')
            try:
                data = await asyncio.to_thread(pdf_page, data, page)
            except Exception as exc:
                raise HTTPException(400, service.safe_error(exc)) from exc
            return Response(data, media_type='image/png', headers={'Cache-Control': 'no-store'})
        image_type = {'.png': 'image/png', '.jpg': 'image/jpeg', '.jpeg': 'image/jpeg', '.gif': 'image/gif', '.webp': 'image/webp'}.get(Path(path).suffix.lower())
        return Response(data, media_type=image_type or 'application/octet-stream', headers={
            'Content-Disposition': ('inline' if image_type else 'attachment')+"; filename*=UTF-8''"+quote(Path(path).name),
            'X-Content-Type-Options': 'nosniff', 'Cache-Control': 'no-store'})

    @app.get('/api/sessions/{sid}/file-diff')
    async def file_diff(sid: str, turn: str, path: str):
        require(sid)
        try:
            return await asyncio.to_thread(versions.diff, sid, turn, path)
        except (ValueError, OSError) as exc:
            raise HTTPException(409, str(exc)) from exc

    @app.post('/api/sessions/{sid}/undo')
    async def undo(sid: str, body: UndoRequest, request: Request):
        session = require(sid)
        origin = request.headers.get('origin')
        if origin and urlparse(origin).netloc != request.headers.get('host'):
            raise HTTPException(403, '仅支持当前工作区发起的撤销')
        from .workspace_actions import WorkspaceActions
        actions = WorkspaceActions(service)
        if actions.metadata('session', sid).get('archived') or actions.metadata('project', session['project_id']).get('archived'):
            raise HTTPException(409, '请先恢复归档对话')
        lock = service.locks.setdefault(sid, asyncio.Lock())
        run = service.runs.db.latest(sid)
        if lock.locked() or (run and run['status'] in {'queued', 'running', 'recovering', 'waiting_input'}):
            raise HTTPException(409, '任务正在运行或等待回答；请先停止任务，再撤销文件')
        async with lock:
            try:
                result = await asyncio.to_thread(versions.undo, sid, body.turn_id, body.paths)
                from langchain_core.messages import SystemMessage
                store.append(sid, SystemMessage(content='工作区文件操作：用户通过界面撤销了文件更改：'+ '、'.join(result['restored'])+'。这些路径已恢复至修改前版本；后续使用前请重新读取文件，不要沿用对话中的旧内容。'))
                store.event(sid, 'ui_file_undo', {'turn_id': body.turn_id, 'paths': result['restored']})
                timeline.emit(sid, body.turn_id, 'files', timeline.files(sid, body.turn_id))
                timeline.emit(sid, body.turn_id, 'note', {'text': '已撤销文件更改：'+ '、'.join(result['restored'])})
                return result
            except (ValueError, OSError) as exc:
                raise HTTPException(409, str(exc)) from exc
